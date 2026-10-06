import io
from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}


class EmptyDocumentError(ValueError):
    """The file was readable but contained no extractable text."""


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF is password-protected.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".txt": _extract_txt,
    ".docx": _extract_docx,
}


def extract_text(ext: str, data: bytes) -> str:
    text = _EXTRACTORS[ext](data).strip()
    if not text:
        raise EmptyDocumentError(
            "No text found. Scanned/image-only PDFs need OCR, which isn't supported."
        )
    return text